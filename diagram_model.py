"""Document model (plain dicts saved as JSON) plus export helpers.

doc = {"app": "NetworkDiagram", "version": 1, "next_id": int,
       "images": {name: base64 png},
       "pages": [{"name": str, "shapes": [shape...], "conns": [connector...]}]}
shape = {"id", "kind", "x", "y", "w", "h", "label", "fill", "line", "text_color", "font",
         "bold", "lw", "show_ip", "data": {"ip", "model", "serial", "location", "notes"}}
conn  = {"id", "src", "dst", "style", "color", "width", "dash", "arrow",
         "label", "src_label", "dst_label", "font"}
Shapes are kept in z-order; containers always draw underneath everything else.
"""
import copy
import json

from stencils import STENCILS
import render

DATA_FIELDS = ("ip", "model", "serial", "location", "notes")
APP_TAG = "NetworkDiagram"


def new_doc():
    return {"app": APP_TAG, "version": 1, "next_id": 1, "images": {},
            "pages": [new_page("Page-1")]}


def new_page(name):
    return {"name": name, "shapes": [], "conns": []}


def next_id(doc, prefix):
    n = doc.get("next_id", 1)
    doc["next_id"] = n + 1
    return "%s%d" % (prefix, n)


def default_label(doc, kind):
    st = STENCILS.get(kind)
    if kind == "text":
        return "Text"
    if kind == "container":
        return "Zone"
    if not st or not st["prefix"]:
        return st["name"] if st and st["label"] == "below" else ""
    used = {s.get("label", "") for p in doc["pages"] for s in p["shapes"]}
    n = 1
    while "%s%02d" % (st["prefix"], n) in used:
        n += 1
    return "%s%02d" % (st["prefix"], n)


def new_shape(doc, kind, cx, cy, w=None, h=None):
    st = STENCILS.get(kind)
    if w is None:
        w, h = (st["w"], st["h"]) if st else (64, 64)
    s = {"id": next_id(doc, "s"), "kind": kind, "x": cx - w / 2, "y": cy - h / 2, "w": w, "h": h,
         "label": default_label(doc, kind), "font": 13 if kind == "container" else 11,
         "show_ip": True, "data": {f: "" for f in DATA_FIELDS}}
    if kind == "text":
        s["font"] = 14
    return s


def new_conn(doc, src, dst, style="straight"):
    return {"id": next_id(doc, "c"), "src": src, "dst": dst, "style": style,
            "color": "#2c3e50", "width": 2, "dash": "solid", "arrow": "none",
            "label": "", "src_label": "", "dst_label": "", "font": 9}


def is_container(s):
    return s["kind"] == "container"


def draw_order(page):
    return [s for s in page["shapes"] if is_container(s)] + \
           [s for s in page["shapes"] if not is_container(s)]


def load(path):
    with open(path, "r", encoding="utf-8") as f:
        doc = json.load(f)
    if doc.get("app") != APP_TAG or not doc.get("pages"):
        raise ValueError("Not a Network Diagram file.")
    doc.setdefault("images", {})
    max_id = 0
    for p in doc["pages"]:
        p.setdefault("shapes", [])
        p.setdefault("conns", [])
        ids = {s["id"] for s in p["shapes"]}
        p["conns"] = [c for c in p["conns"] if c["src"] in ids and c["dst"] in ids]
        for item in p["shapes"] + p["conns"]:
            digits = "".join(ch for ch in item["id"] if ch.isdigit())
            max_id = max(max_id, int(digits or 0))
        for s in p["shapes"]:
            s.setdefault("data", {})
            for f in DATA_FIELDS:
                s["data"].setdefault(f, "")
    doc["next_id"] = max(doc.get("next_id", 1), max_id + 1)
    return doc


def save(doc, path):
    used = {s["kind"][4:] for p in doc["pages"] for s in p["shapes"] if s["kind"].startswith("img:")}
    out = dict(doc)
    out["images"] = {k: v for k, v in doc.get("images", {}).items() if k in used}
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(out, f, indent=1)
    import os
    os.replace(tmp, path)


def clone_items(doc, shapes, conns, dx=0, dy=0):
    """Deep-copy shapes (and the connectors between them) with fresh ids."""
    idmap = {}
    new_shapes, new_conns = [], []
    for s in shapes:
        n = copy.deepcopy(s)
        n["id"] = next_id(doc, "s")
        n["x"] += dx
        n["y"] += dy
        idmap[s["id"]] = n["id"]
        new_shapes.append(n)
    for c in conns:
        if c["src"] in idmap and c["dst"] in idmap:
            n = copy.deepcopy(c)
            n["id"] = next_id(doc, "c")
            n["src"], n["dst"] = idmap[c["src"]], idmap[c["dst"]]
            new_conns.append(n)
    return new_shapes, new_conns


# --- geometry over a page -----------------------------------------------------------

def connector_paths(page):
    by_id = {s["id"]: s for s in page["shapes"]}
    offs = render.parallel_offsets(page["conns"])
    out = {}
    for c in page["conns"]:
        a, b = by_id.get(c["src"]), by_id.get(c["dst"])
        if a and b:
            out[c["id"]] = render.route(c, a, b, offs.get(c["id"], 0))
    return out


def page_bbox(page, paths=None):
    xs0, ys0, xs1, ys1 = [], [], [], []
    for s in page["shapes"]:
        x0, y0, x1, y1 = render.shape_bbox(s)
        xs0.append(x0); ys0.append(y0); xs1.append(x1); ys1.append(y1)
    for pts in (paths or connector_paths(page)).values():
        for x, y in pts:
            xs0.append(x); ys0.append(y); xs1.append(x); ys1.append(y)
    if not xs0:
        return None
    return min(xs0), min(ys0), max(xs1), max(ys1)


def render_page(backend, page, bg="#ffffff"):
    paths = connector_paths(page)
    for s in draw_order(page):
        if is_container(s):
            render.draw_shape(backend, s)
    for c in page["conns"]:
        if c["id"] in paths:
            render.draw_connector(backend, c, paths[c["id"]], bg)
    for s in draw_order(page):
        if not is_container(s):
            render.draw_shape(backend, s)


# --- exports -----------------------------------------------------------------------------

def page_to_image(page, images, scale=2.0, margin=30, bg="#ffffff"):
    """Anti-aliased PIL image of a page (rendered at 2x then downsampled)."""
    from PIL import Image
    bb = page_bbox(page)
    if not bb:
        raise ValueError("The page is empty.")
    x0, y0, x1, y1 = bb
    ss = scale * 2
    W = int((x1 - x0 + 2 * margin) * ss)
    H = int((y1 - y0 + 2 * margin) * ss)
    if W * H > 120_000_000:
        raise ValueError("Diagram is too large to export at this resolution.")
    img = Image.new("RGB", (W, H), bg)
    b = render.PilBackend(img, ss, (margin - x0) * ss, (margin - y0) * ss, images)
    render_page(b, page, bg)
    return img.resize((W // 2, H // 2), Image.LANCZOS)


def page_to_svg(page, images, margin=30, bg="#ffffff"):
    bb = page_bbox(page)
    if not bb:
        raise ValueError("The page is empty.")
    x0, y0, x1, y1 = bb
    b = render.SvgBackend(1.0, margin - x0, margin - y0, images)
    render_page(b, page, bg)
    return b.document(int(x1 - x0 + 2 * margin), int(y1 - y0 + 2 * margin), bg)


def inventory_rows(doc):
    devices, links = [], []
    for p in doc["pages"]:
        by_id = {s["id"]: s for s in p["shapes"]}
        for s in p["shapes"]:
            if s["kind"] in ("container", "text") or s["kind"] in ("rect", "rounded", "ellipse",
                                                                    "diamond", "hexagon"):
                continue
            st = STENCILS.get(s["kind"])
            kind = st["name"] if st else s["kind"][4:].split("/")[-1]
            d = s.get("data", {})
            devices.append([p["name"], s.get("label", "").replace("\n", " "), kind, d.get("ip", ""),
                            d.get("model", ""), d.get("serial", ""), d.get("location", ""),
                            d.get("notes", "")])
        for c in p["conns"]:
            a, b = by_id.get(c["src"]), by_id.get(c["dst"])
            if not a or not b:
                continue
            links.append([p["name"], a.get("label", "").replace("\n", " "), c.get("src_label", ""),
                          b.get("label", "").replace("\n", " "), c.get("dst_label", ""),
                          c.get("label", "").replace("\n", " ")])
    return devices, links


DEVICE_HEAD = ["Page", "Name", "Type", "IP Address", "Model", "Serial", "Location", "Notes"]
LINK_HEAD = ["Page", "From Device", "From Interface", "To Device", "To Interface", "Link Label"]


def export_inventory(doc, path):
    devices, links = inventory_rows(doc)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    for idx, (title, head, rows) in enumerate((("Devices", DEVICE_HEAD, devices),
                                                ("Links", LINK_HEAD, links))):
        ws = wb.active if idx == 0 else wb.create_sheet()
        ws.title = title
        ws.append(head)
        for r in rows:
            ws.append(r)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4F81BD")
            cell.alignment = Alignment(vertical="center")
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(50, max(10, width + 2))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)
    return len(devices), len(links)
